from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from investigation_engine.application.chunking import ChunkingConfig, ChunkingStrategy, chunk_document
from investigation_engine.application.ports.text_tokenizer import TokenizerIdentity
from investigation_engine.application.use_cases.register_document import RegisterDocument
from investigation_engine.domain.source_span import LocatorKind
from investigation_engine.infrastructure.ingestion.docx_adapter import extract_docx
from investigation_engine.infrastructure.ingestion.pdf_adapter import extract_pdf
from investigation_engine.infrastructure.ingestion.extraction_errors import (
    DocumentExtractionError, InsufficientTextError, UnsupportedDocumentError,
)
from investigation_engine.infrastructure.persistence.in_memory_document_repository import InMemoryDocumentRepository


FIXTURES = Path("data/samples/formats")


def test_pdf_real_multipage_unicode_empty_page_metadata_and_exact_offsets():
    document = extract_pdf(FIXTURES / "multipage.pdf", source_id="independent-upload")
    assert document.title == "Field observations"
    assert "Café observations: naïve, Ω and résumé." in document.text
    assert [(s.start, s.end, s.locators[0].kind, s.locators[0].index) for s in document.source_spans] == [
        (0, 82, LocatorKind.PAGE, 1), (84, 157, LocatorKind.PAGE, 3)]
    assert document.text[82:84] == "\n\n"
    assert document.text[84:].startswith("Third page")
    assert document.source.source_id == "independent-upload"
    assert document.source.metadata["empty_pages"] == [2]
    assert document.source.metadata["page_count"] == 3
    assert document.source.metadata["pdf"]["author"] == "Fixture author"
    assert all(type(v) is str for v in document.source.metadata["pdf"].values())


def test_pdf_without_metadata_uses_filename_and_preserves_text(tmp_path):
    path = tmp_path / "untitled.pdf"
    writer = PdfWriter(clone_from=FIXTURES / "multipage.pdf")
    writer.metadata = None
    writer.write(path)
    extracted = extract_pdf(path, source_id="new-upload")
    assert extracted.title == path.name and extracted.source.metadata["pdf"] == {}
    assert extracted.text == extract_pdf(FIXTURES / "multipage.pdf", source_id="original").text


def test_pdf_without_extractable_text_and_encrypted_files_have_explicit_outcomes(tmp_path):
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    with pytest.raises(InsufficientTextError, match="OCR"):
        extract_pdf(path, source_id="blank")
    writer.encrypt("password", algorithm="RC4-128")
    writer.write(path)
    with pytest.raises(UnsupportedDocumentError, match="encrypted"):
        extract_pdf(path, source_id="encrypted")


def test_pdf_page_failure_rejects_partial_extraction_and_closes_stream(monkeypatch):
    from investigation_engine.infrastructure.ingestion import pdf_adapter
    streams = []

    def fail(**kwargs):
        raise ValueError("bad page content")

    def reader(stream, **kwargs):
        streams.append(stream)
        return SimpleNamespace(is_encrypted=False, metadata={}, pages=[
            SimpleNamespace(extract_text=lambda **kwargs: "First readable page"),
            SimpleNamespace(extract_text=fail)])

    monkeypatch.setattr(pdf_adapter.pypdf, "PdfReader", reader)
    with pytest.raises(DocumentExtractionError):
        extract_pdf(FIXTURES / "multipage.pdf", source_id="broken-page")
    assert streams[0].closed


def test_docx_paragraph_table_reading_order_unicode_metadata_and_structural_coordinates():
    document = extract_docx(FIXTURES / "paragraphs-and-table.docx", source_id="upload")
    assert document.text == ("Café observations: naïve, Ω and résumé. 😀\n\n"
                             "Station\tReading\nSeoul 서울\t23 °C\nAfter the table: astronomy archive.")
    assert [document.text[s.start:s.end] for s in document.source_spans] == [
        "Café observations: naïve, Ω and résumé. 😀", "Station", "Reading", "Seoul 서울", "23 °C",
        "After the table: astronomy archive."]
    locators = document.source_spans[3].locators
    assert [(loc.kind.value, loc.index) for loc in locators] == [("table", 0), ("row", 1), ("cell", 0), ("paragraph", 0)]
    assert document.source_spans[-1].locators[0].index == 2  # Count the original empty paragraph.
    assert all(loc.kind != LocatorKind.PAGE for span in document.source_spans for loc in span.locators)
    assert document.source.metadata["core_properties"]["author"] == "Fixture author"
    assert type(document.source.metadata["core_properties"]["revision"]) is int
    assert document.source.metadata["format"] == "docx"


def test_docx_nested_tables_and_merged_cells_do_not_repeat_text(tmp_path):
    path = tmp_path / "nested.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged source"
    table.cell(1, 0).text = "Before"
    table.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested source"
    table.cell(1, 1).text = "Last"
    document.save(path)
    extracted = extract_docx(path, source_id="nested")
    assert extracted.text.count("Merged source") == 1
    assert extracted.text.index("Before") < extracted.text.index("Nested source") < extracted.text.index("Last")
    nested = next(span for span in extracted.source_spans if extracted.text[span.start:span.end] == "Nested source")
    assert [loc.kind.value for loc in nested.locators] == ["table", "row", "cell", "table", "row", "cell", "paragraph"]
    assert extracted.title == path.name


def test_empty_docx_has_explicit_insufficient_text_outcome(tmp_path):
    path = tmp_path / "empty.docx"
    document = Document()
    document.add_paragraph(" \t")
    document.save(path)
    with pytest.raises(InsufficientTextError, match="body text"):
        extract_docx(path, source_id="empty")


@pytest.mark.parametrize("extractor,extension", [(extract_pdf, "pdf"), (extract_docx, "docx")])
@pytest.mark.parametrize("body", [b"", b"not a document", b"%PDF-1.7\ntruncated"])
def test_invalid_files_fail_without_substituting_content(tmp_path, extractor, extension, body):
    path = tmp_path / f"invalid.{extension}"
    path.write_bytes(body)
    with pytest.raises(DocumentExtractionError):
        extractor(path, source_id="invalid")


def test_valid_zip_that_is_not_docx_fails_safely(tmp_path):
    path = tmp_path / "not-word.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("other.txt", "unrelated")
    with pytest.raises(DocumentExtractionError):
        extract_docx(path, source_id="wrong-package")


def test_docx_missing_core_metadata_is_not_replaced_by_library_defaults(tmp_path):
    path = tmp_path / "no-properties.docx"
    with ZipFile(FIXTURES / "paragraphs-and-table.docx") as original, ZipFile(path, "w") as changed:
        for name in original.namelist():
            if name == "docProps/core.xml":
                continue
            body = original.read(name)
            if name in ("_rels/.rels", "[Content_Types].xml"):
                root = ElementTree.fromstring(body)
                for child in list(root):
                    if child.attrib.get("Target") == "docProps/core.xml" or child.attrib.get("PartName") == "/docProps/core.xml":
                        root.remove(child)
                body = ElementTree.tostring(root)
            changed.writestr(name, body)
    extracted = extract_docx(path, source_id="absent-properties")
    assert extracted.title == path.name and extracted.source.metadata["core_properties"] == {}
    assert extracted.text == extract_docx(FIXTURES / "paragraphs-and-table.docx", source_id="original").text


def test_docx_malformed_main_xml_fails_without_partial_result(tmp_path):
    path = tmp_path / "broken-xml.docx"
    with ZipFile(FIXTURES / "paragraphs-and-table.docx") as original, ZipFile(path, "w") as changed:
        for name in original.namelist():
            changed.writestr(name, b"<unclosed>" if name == "word/document.xml" else original.read(name))
    with pytest.raises(DocumentExtractionError):
        extract_docx(path, source_id="broken-xml")


@pytest.mark.parametrize("extractor,filename", [(extract_pdf, "multipage.pdf"), (extract_docx, "paragraphs-and-table.docx")])
async def test_registration_retries_content_reuse_and_chunking_use_existing_interfaces(extractor, filename):
    extracted = extractor(FIXTURES / filename, source_id="first-upload")
    register = RegisterDocument(InMemoryDocumentRepository())
    first, version = await register.execute(extracted, owner_id="principal")
    retry, same = await register.execute(extracted, owner_id="principal")
    independent, other = await register.execute(replace(extracted, source=replace(extracted.source, source_id="second-upload")), owner_id="principal")
    assert first.id == retry.id and version.id == same.id
    assert independent.id != first.id and other.content is version.content

    class Tokenizer:
        identity = TokenizerIdentity(name="characters", revision="1", artifact_sha256="0" * 64, runtime_version="test")

        def offsets(self, text):
            return tuple((i, i + 1) for i in range(len(text)))

    config = ChunkingConfig(strategy=ChunkingStrategy.FIXED_TOKENS, tokenizer=Tokenizer.identity, max_tokens=30)
    plan = chunk_document(version, config=config, tokenizer=Tokenizer())
    assert "".join(c.text for c in plan.chunks) == extracted.text


@pytest.mark.parametrize("extractor", [extract_pdf, extract_docx])
def test_source_ids_are_caller_owned_and_missing_files_are_not_hidden(tmp_path, extractor):
    with pytest.raises(ValueError, match="Source ID"):
        extractor(tmp_path / "missing", source_id=" ")
    with pytest.raises(FileNotFoundError):
        extractor(tmp_path / "missing", source_id="stable")
